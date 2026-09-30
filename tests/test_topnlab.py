"""TopNLab integration. ТЗ «Интеграция с TopNLab CRM» v1.0.

The HTTP side runs against httpx.MockTransport answering the way TopNLab's docs
say -- and, where they differ, the way the live API was seen to answer
(menu/get-all-pages returns ``response``, not ``data``). Task types below are
the real ones of the customer's account, fetched read-only on 30.09.2026.
"""
import json
import os
from datetime import datetime, timezone
from types import SimpleNamespace

import httpx
import pytest

from app.services import topnlab_adapter as tl

pytestmark_db = pytest.mark.skipif(os.getenv("RUN_DB_TESTS") != "1",
                                   reason="requires live PostgreSQL")

APPKEY = "05ab9dc0-test-appkey-000000"

LIVE_TASK_TYPES = [
    {"id": 35307, "name": "Позвонить", "custom_task_type_cards": [1, 2, 3, 4, 196932, 196933],
     "is_hidden": False},
    {"id": 35308, "name": "Встреча в офисе", "custom_task_type_cards": [1, 2, 3, 4],
     "is_hidden": False},
    {"id": 160072, "name": "Напоминание)", "custom_task_type_cards": [1, 3], "is_hidden": True},
    {"id": 230550, "name": "МЛС. Предложите объекты", "custom_task_type_cards": [3],
     "is_hidden": False, "is_note_required": True},
]


def _settings(**kw):
    return tl.TopnlabSettings(agency_id="agency-1", appkey=APPKEY, **kw)


def _lead(**kw):
    base = dict(
        id="11111111-1111-4111-8111-111111111111", agency_id="agency-1",
        name=None, phone=None, telegram_username="buyer_gel", intent_score=82,
        segment="family", urgency="hot", budget_min=5_000_000, budget_max=8_000_000,
        source_type="signal", buyer_profile={}, consent_given=True,
        topnlab_client_id=None, topnlab_task_id=None, topnlab_synced_at=None,
        crm_deal_id=None, source_signal_id="sig-1", signal_id="sig-1",
    )
    base.update(kw)
    return SimpleNamespace(**base)


def _signal(**kw):
    base = dict(
        raw_text="Ищу 2-комнатную в Геленджике у моря, бюджет до 8 млн, ипотека одобрена",
        signal_url="https://t.me/gelendzhik_chat/4512",
        created_at=datetime(2026, 9, 30, 9, 0, tzinfo=timezone.utc),
        source=SimpleNamespace(source_name="Геленджик Недвижимость", source_type="telegram_chat"),
        ai_analysis={"property_type": "apartment"},
    )
    base.update(kw)
    return SimpleNamespace(**base)


class _Api:
    """A fake TopNLab: routes by path, records every request."""

    def __init__(self, routes):
        self.routes = routes
        self.calls = []

    def handler(self, request: httpx.Request) -> httpx.Response:
        body = json.loads(request.content) if request.content else None
        self.calls.append((request.method, request.url.path, dict(request.url.params), body))
        answer = self.routes.get(request.url.path)
        if answer is None:
            return httpx.Response(404, json={"status": "error"})
        if callable(answer):
            return answer(request)
        status, payload = answer
        return httpx.Response(status, json=payload)

    def paths(self):
        return [c[1] for c in self.calls]


@pytest.fixture
def api(monkeypatch):
    """Install a fake TopNLab and an in-memory task-types cache."""
    real_client = httpx.AsyncClient
    holder = {}

    def install(routes):
        fake = _Api(routes)
        holder["api"] = fake
        monkeypatch.setattr(tl.httpx, "AsyncClient",
                            lambda **kw: real_client(transport=httpx.MockTransport(fake.handler), **kw))
        return fake

    cache = {}

    async def cache_get(key):
        return cache.get(key)

    async def cache_set(key, value):
        cache[key] = value

    monkeypatch.setattr(tl, "_cache_get", cache_get)
    monkeypatch.setattr(tl, "_cache_set", cache_set)
    install.cache = cache
    return install


# --------------------------------------------------------------- payload shape

def test_comment_carries_score_segment_budget_source_and_contact():
    comment = tl.build_comment(_lead(), _signal())
    assert comment.startswith("[REIP] Score: 82/100")
    assert "Сегмент: семья" in comment and "Срочность: горячий" in comment
    assert "Бюджет: 5000–8000к ₽" in comment
    assert "Источник: Геленджик Недвижимость (telegram_chat)" in comment
    # Without these the manager in TopNLab has no way to reach the person.
    assert "Telegram: @buyer_gel" in comment
    assert "https://t.me/gelendzhik_chat/4512" in comment
    assert "ипотека одобрена" in comment


def test_comment_is_cut_to_500_from_the_message_end_not_the_contact():
    comment = tl.build_comment(_lead(), _signal(raw_text="очень длинно " * 200))
    assert len(comment) <= 500
    assert "https://t.me/gelendzhik_chat/4512" in comment
    assert comment.endswith("…»")


def test_comment_without_budget_or_signal():
    comment = tl.build_comment(_lead(budget_min=None, budget_max=None, segment=None,
                                     urgency=None, telegram_username=None))
    assert comment == "[REIP] Score: 82/100 | Источник: signal"


@pytest.mark.parametrize("value,expected", [
    ("apartment", "flat"), ("studio", "flat"), ("house", "house"), ("townhouse", "house"),
    ("land", "land"), ("commercial", "commerce"), ("garage", "garage"), (None, "flat"),
])
def test_object_type_maps_to_topnlab_vocabulary(value, expected):
    assert tl.object_type_for(_lead(buyer_profile={"property_type": value})) == expected


def test_client_payload_defaults_and_virtual_number():
    payload = tl.build_client_payload(_settings(virtual_number="8612345678"), _lead(), _signal())
    assert payload["appkey"] == APPKEY
    assert payload["fullname"] == "Покупатель из REIP"
    assert payload["phone"] == ""
    assert payload["action"] == 1
    assert payload["object_type"] == "flat"
    assert payload["to_number"] == "8612345678"

    named = tl.build_client_payload(_settings(), _lead(name="Анна", phone="+7 (918) 000-11-22"))
    assert named["fullname"] == "Анна" and named["phone"] == "79180001122"
    assert "to_number" not in named


def test_the_call_task_type_is_found_among_the_live_types():
    assert tl.find_call_task_type(LIVE_TASK_TYPES) == 35307
    hidden_only = [{"id": 9, "name": "Позвонить", "is_hidden": True}]
    assert tl.find_call_task_type(hidden_only) == 9  # better than no task at all
    assert tl.find_call_task_type([{"id": 1, "name": "Встреча"}]) is None


# ------------------------------------------------------- answers per the docs

@pytest.mark.asyncio
async def test_import_client_returns_inserted_id(api):
    fake = api({"/call/main/importClient/": (200, {"status": "ok", "insertedId": 100500})})
    assert await tl.push_lead_to_topnlab(_settings(), _lead(), _signal()) == 100500
    method, path, _, body = fake.calls[0]
    assert (method, path) == ("POST", "/call/main/importClient/")
    assert body["appkey"] == APPKEY and body["action"] == 1


@pytest.mark.asyncio
async def test_success_with_warnings_is_still_success(api):
    api({"/call/main/importClient/": (200, {"status": "ok", "insertedId": 7,
                                             "errors": {"object_id": "нет такого"}})})
    assert await tl.push_lead_to_topnlab(_settings(), _lead()) == 7


@pytest.mark.asyncio
async def test_422_is_rejected_and_not_retried(api):
    api({"/call/main/importClient/": (422, {"status": "error",
                                             "errors": {"action": "Не указан тип сделки"}})})
    with pytest.raises(tl.TopnlabRejected) as err:
        await tl.push_lead_to_topnlab(_settings(), _lead())
    assert not isinstance(err.value, tl.TopnlabUnavailable)
    assert err.value.errors == {"action": "Не указан тип сделки"}


@pytest.mark.asyncio
@pytest.mark.parametrize("answer", [
    lambda r: httpx.Response(500, text="Internal error"),
    lambda r: httpx.Response(200, text="<html>maintenance</html>"),
    lambda r: httpx.Response(502, json={"status": "ok"}),
])
async def test_anything_undocumented_is_temporary(api, answer):
    api({"/call/main/importClient/": answer})
    with pytest.raises(tl.TopnlabUnavailable):
        await tl.push_lead_to_topnlab(_settings(), _lead())


@pytest.mark.asyncio
async def test_timeout_is_temporary_and_the_key_stays_out_of_logs(api, monkeypatch):
    def boom(request):
        raise httpx.ConnectTimeout(f"timed out: {request.url}", request=request)

    logged = []
    monkeypatch.setattr(tl.logger, "warning", lambda event, **kw: logged.append((event, kw)))
    api({"/api/partners/tasks/get-task-types": boom})
    with pytest.raises(tl.TopnlabUnavailable):
        await tl.get_topnlab_task_types(_settings(), use_cache=False)
    assert logged and APPKEY not in json.dumps(logged, ensure_ascii=False)


@pytest.mark.asyncio
async def test_ok_without_inserted_id_is_not_retried(api):
    """A retry could create a second order; without an id there is nothing to
    hang the task on either."""
    api({"/call/main/importClient/": (200, {"status": "ok"})})
    with pytest.raises(tl.TopnlabRejected):
        await tl.push_lead_to_topnlab(_settings(), _lead())


# ------------------------------------------------------------- task «Позвонить»

@pytest.mark.asyncio
async def test_call_task_goes_on_the_order_in_moscow_time(api):
    fake = api({
        "/api/partners/tasks/get-task-types": (200, {"status": "success", "data": LIVE_TASK_TYPES}),
        "/api/partners/tasks/create": (200, {"status": "success", "data": [{"id": 6104297}]}),
    })
    now = datetime(2026, 9, 30, 9, 0, tzinfo=timezone.utc)
    task_id = await tl.create_topnlab_task(_settings(), 100500, _lead(), 5, now=now)
    assert task_id == 6104297
    _, path, params, body = fake.calls[0]
    assert path == "/api/partners/tasks/get-task-types" and params["appkey"] == APPKEY
    body = fake.calls[1][3]
    assert body["owner_id"] == 100500 and body["owner_type"] == 2
    assert body["task_type_id"] == 35307
    assert body["begin_at"] == "2026-09-30 12:05:00"  # 09:00 UTC + 5 min, MSK
    assert body["end_at"] == "2026-09-30 12:35:00"
    assert body["notify_before_value"] == "5" and body["notify_before_type"] == 2
    assert body["description"].startswith("Горячий лид из REIP! Score: 82/100")


@pytest.mark.asyncio
async def test_task_types_are_cached_for_a_day(api):
    fake = api({"/api/partners/tasks/get-task-types":
                (200, {"status": "success", "data": LIVE_TASK_TYPES})})
    await tl.get_topnlab_task_types(_settings())
    await tl.get_topnlab_task_types(_settings())
    assert fake.paths().count("/api/partners/tasks/get-task-types") == 1
    assert "topnlab_task_types:agency-1" in api.cache


@pytest.mark.asyncio
async def test_no_call_type_means_no_task_and_no_error(api):
    fake = api({"/api/partners/tasks/get-task-types":
                (200, {"status": "success", "data": [{"id": 1, "name": "Встреча"}]})})
    assert await tl.create_topnlab_task(_settings(), 1, _lead()) is None
    assert "/api/partners/tasks/create" not in fake.paths()


# -------------------------------------------------- call log, manager, Avito

@pytest.mark.asyncio
async def test_call_log_needs_the_virtual_number_and_never_raises(api):
    fake = api({"/call/main/importCall/": (500, {})})
    assert await tl.log_signal_as_call(_settings(), 1, _signal()) is False
    assert fake.calls == []
    assert await tl.log_signal_as_call(_settings(virtual_number="8612345678"), 1, _signal()) is False

    fake = api({"/call/main/importCall/": (200, {"status": "ok"})})
    assert await tl.log_signal_as_call(_settings(virtual_number="8612345678"), 77, _signal())
    body = fake.calls[0][3]
    assert body["direction"] == "income" and body["answered"] == 0
    assert body["bind_to_client_id"] == 77 and body["to_number"] == "8612345678"
    assert body["end_time"] - body["start_time"] == 60


@pytest.mark.asyncio
async def test_transfer_only_with_an_email(api):
    fake = api({"/call/main/transferClient/": (200, {"status": "ok"})})
    assert await tl.transfer_to_manager(_settings(), 5) is False
    assert await tl.transfer_to_manager(_settings(manager_email="agent@agency.ru"), 5)
    assert fake.calls[0][3] == {"appkey": APPKEY, "client_id": 5, "user_mail": "agent@agency.ru"}


@pytest.mark.asyncio
async def test_avito_credentials_need_the_company_id(api):
    fake = api({"/public/partner/207413/credentials":
                (200, {"status": "ok", "data": {"client_id": "a", "client_secret": "b"}})})
    with pytest.raises(tl.TopnlabRejected):
        await tl.get_avito_credentials(_settings())
    creds = await tl.get_avito_credentials(_settings(company_id="207413"))
    assert creds == {"client_id": "a", "client_secret": "b"}
    assert fake.calls[0][2] == {"site": "AVITO", "key": APPKEY}


def test_credentials_are_stored_encrypted():
    blob = tl.encrypt_blob({"client_secret": "very-secret"})
    assert "very-secret" not in blob
    assert tl.decrypt_blob(blob) == {"client_secret": "very-secret"}


# ------------------------------------------------------------- report button

PAGES = {"status": "ok", "response": {"7": "Заявки - Аренда", "8": "Заявки - Продажа",
                                      "196933": "Заявки на Ипотеку"}}


@pytest.mark.asyncio
async def test_report_is_registered_on_the_sale_orders_page(api):
    fake = api({
        "/public/menu/list/": (200, {"status": "ok", "data": []}),
        "/public/menu/get-all-pages": (200, PAGES),
        "/public/menu/create": (200, {"status": "ok", "data": {"id": 888}}),
    })
    menu_id, token, created = await tl.register_reip_report(_settings())
    assert (menu_id, created) == (888, True)
    assert len(token) >= 24
    params = [c for c in fake.calls if c[1] == "/public/menu/create"][0][2]
    assert params["page_id"] == "8" and params["title"] == "Аналитика REIP"
    assert params["url"] == f"https://test.local/api/topnlab/report-webhook?token={token}"


@pytest.mark.asyncio
async def test_registering_twice_keeps_one_button(api):
    url = tl.report_webhook_url("tok-existing-000000000000")
    fake = api({"/public/menu/list/": (200, {"status": "ok", "data": [
        {"id": 3, "title": "Другое", "url": "https://x"},
        {"id": 888, "title": "Аналитика REIP", "url": url}]})})
    menu_id, token, created = await tl.register_reip_report(
        _settings(report_token="tok-existing-000000000000"))
    assert (menu_id, token, created) == (888, "tok-existing-000000000000", False)
    assert "/public/menu/create" not in fake.paths()


def test_report_workbook():
    from app.routers.topnlab_webhooks import HEADERS, build_report_xlsx

    data = build_report_xlsx([[100500, "30.09.2026 09:00", 82, "семья", "горячий", 5_000_000,
                               8_000_000, "new", "", "@buyer", "чат", "https://t.me/x/1", "текст"]], 2)
    assert data[:2] == b"PK"
    from io import BytesIO

    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(data)).active
    assert [c.value for c in ws[1]] == HEADERS
    assert ws["A2"].value == 100500
    empty = load_workbook(BytesIO(build_report_xlsx([], 3))).active
    assert "нет пришедших из REIP" in empty["A2"].value


# ----------------------------------------------------------------- the flow

class _Session:
    def __init__(self, cfg, signal=None):
        self.cfg, self.signal, self.commits = cfg, signal, 0

    async def get(self, model, pk):
        return self.signal

    async def commit(self):
        self.commits += 1


def _cfg(**extra):
    return SimpleNamespace(agency_id="agency-1", crm_type="topnlab", api_key=APPKEY,
                           base_url=None, is_active=True,
                           config={"sync_enabled": True, **extra})


@pytest.fixture
def flow(monkeypatch, api):
    from app.config import config

    monkeypatch.setattr(config, "topnlab_sync_enabled", True)

    def run(cfg, routes, lead, signal=None, **kw):
        fake = api(routes)

        async def load(session, agency_id):
            return cfg

        monkeypatch.setattr(tl, "load_config", load)
        session = _Session(cfg, signal)
        return fake, session, tl.sync_lead_to_topnlab(session, lead, **kw)

    return run


ALL_OK = {
    "/call/main/importClient/": (200, {"status": "ok", "insertedId": 100500}),
    "/api/partners/tasks/get-task-types": (200, {"status": "success", "data": LIVE_TASK_TYPES}),
    "/api/partners/tasks/create": (200, {"status": "success", "data": [{"id": 6104297}]}),
    "/call/main/importCall/": (200, {"status": "ok"}),
    "/call/main/transferClient/": (200, {"status": "ok"}),
}


@pytest.mark.asyncio
async def test_hot_lead_becomes_order_with_call_task(flow):
    lead = _lead()
    cfg = _cfg(virtual_number="8612345678", manager_email="agent@agency.ru")
    fake, session, run = flow(cfg, ALL_OK, lead, _signal())
    result = await run
    assert result["exported"] is True
    assert lead.topnlab_client_id == 100500 and lead.topnlab_task_id == 6104297
    assert lead.topnlab_synced_at is not None
    assert lead.crm_deal_id == "100500"  # attribution chain, as with other CRMs
    assert result["transferred"] is True and result["call_logged"] is True
    assert fake.paths() == ["/call/main/importClient/", "/call/main/transferClient/",
                            "/api/partners/tasks/get-task-types", "/api/partners/tasks/create",
                            "/call/main/importCall/"]


@pytest.mark.asyncio
async def test_a_retry_after_the_order_does_not_create_a_second_one(flow):
    """The order went through, the task call timed out, Celery retried."""
    lead = _lead(topnlab_client_id=100500)
    fake, _, run = flow(_cfg(), ALL_OK, lead, _signal())
    result = await run
    assert "/call/main/importClient/" not in fake.paths()
    assert lead.topnlab_task_id == 6104297 and result["exported"] is True


@pytest.mark.asyncio
async def test_task_timeout_propagates_for_retry_after_the_order_is_saved(flow):
    routes = {**ALL_OK, "/api/partners/tasks/create": lambda r: httpx.Response(503)}
    lead = _lead()
    _, session, run = flow(_cfg(), routes, lead, _signal())
    with pytest.raises(tl.TopnlabUnavailable):
        await run
    assert lead.topnlab_client_id == 100500 and session.commits >= 1


@pytest.mark.asyncio
@pytest.mark.parametrize("change,reason", [
    (dict(global_off=True), "disabled_globally"),
    (dict(cfg=None), "not_configured"),
    (dict(cfg=_cfg(sync_enabled=False)), "disabled_for_agency"),
    (dict(lead=_lead(consent_given=False)), "no_consent"),
    (dict(lead=_lead(intent_score=39)), "below_threshold"),
    (dict(lead=_lead(topnlab_client_id=1, topnlab_task_id=2)), "already_synced"),
])
async def test_what_keeps_a_lead_from_going_out(flow, monkeypatch, change, reason):
    from app.config import config

    lead = change.get("lead", _lead())
    fake, _, run = flow(change.get("cfg", _cfg()), ALL_OK, lead)
    if change.get("global_off"):
        monkeypatch.setattr(config, "topnlab_sync_enabled", False)
    result = await run
    assert result["reason"] == reason
    assert "/call/main/importClient/" not in fake.paths()


@pytest.mark.asyncio
async def test_a_manager_qualifying_by_hand_skips_the_score_gate(flow):
    lead = _lead(intent_score=None)
    _, _, run = flow(_cfg(), ALL_OK, lead, force=True)
    assert (await run)["exported"] is True


@pytest.mark.asyncio
async def test_bad_data_is_reported_not_retried(flow):
    routes = {"/call/main/importClient/": (422, {"status": "error", "errors": {"phone": "?"}})}
    lead = _lead()
    _, _, run = flow(_cfg(), routes, lead)
    result = await run
    assert result == {"exported": False, "reason": "rejected", "errors": {"phone": "?"}}
    assert lead.topnlab_client_id is None


# ------------------------------------------------- existing CRM export path

@pytest.mark.asyncio
async def test_qualified_export_uses_the_topnlab_flow(monkeypatch):
    from app.services import crm_export

    async def cfg(session, agency_id):
        return _cfg()

    seen = {}

    async def fake_sync(session, lead, *, force=False):
        seen["force"] = force
        return {"exported": True, "topnlab_client_id": 1}

    monkeypatch.setattr(crm_export, "_active_config", cfg)
    monkeypatch.setattr(tl, "sync_lead_to_topnlab", fake_sync)
    result = await crm_export.export_lead_to_crm(object(), _lead())
    assert result["exported"] is True and seen["force"] is True


@pytest.mark.asyncio
async def test_a_closed_deal_does_not_open_a_new_topnlab_order(monkeypatch):
    from app.services import crm_export

    async def cfg(session, agency_id):
        return _cfg()

    monkeypatch.setattr(crm_export, "_active_config", cfg)
    result = await crm_export.push_outcome_to_crm(object(), _lead(), SimpleNamespace())
    assert result["reason"] == "not_supported"


def test_the_registry_adapter_points_at_the_real_endpoint():
    from app.services.crm import build_crm_adapter

    a = build_crm_adapter("topnlab", api_key=APPKEY)
    assert a.endpoint() == "https://agencies-p.topnlab.ru/call/main/importClient/"
    assert a.headers() == {}
    assert a.build_payload({"name": None, "phone": "+7 900 1"})["appkey"] == APPKEY


# ------------------------------------------------------------ HTTP endpoints

def _client():
    from fastapi.testclient import TestClient

    from app.main import app

    return TestClient(app)


def test_preflight_from_topnlab_is_answered_despite_the_app_cors_policy():
    resp = _client().options("/api/topnlab/report-webhook", headers={
        "Origin": "https://agencies.topnlab.ru", "Access-Control-Request-Method": "POST"})
    assert resp.status_code == 200
    assert resp.headers["access-control-allow-origin"] == "*"
    assert resp.headers["access-control-allow-headers"] == "*"


def test_other_api_paths_keep_the_app_cors_policy():
    resp = _client().options("/api/leads", headers={
        "Origin": "https://evil.example", "Access-Control-Request-Method": "GET"})
    assert resp.headers.get("access-control-allow-origin") != "*"


def test_report_file_rejects_anything_but_a_generated_name():
    resp = _client().get("/api/topnlab/report-files/..%2F..%2Fsecret.xlsx")
    assert resp.status_code == 404
    resp = _client().get("/api/topnlab/report-files/" + "a" * 32 + ".xlsx")
    assert resp.status_code == 404
    assert resp.headers["access-control-allow-origin"] == "*"


@pytestmark_db
@pytest.mark.asyncio
async def test_report_webhook_end_to_end(tmp_path, monkeypatch):
    """TopNLab posts the chosen orders; REIP answers with a link whose file
    downloads with CORS headers -- the two things ТЗ 5.1 checks."""
    import secrets

    from app.database import async_session, engine, run_migrations
    from app.models.agency import Agency
    from app.models.agency_crm_config import AgencyCRMConfig
    from app.models.lead import Lead
    from app.services import storage

    monkeypatch.setattr(storage, "get_storage", lambda: storage.LocalStorage(tmp_path))
    await run_migrations()
    token = secrets.token_urlsafe(24)
    try:
        async with async_session() as s:
            agency = Agency(name="TopNLab Agency", base_city="Геленджик")
            s.add(agency)
            await s.flush()
            cfg = AgencyCRMConfig(agency_id=agency.id, crm_type="topnlab", is_active=True,
                                  config={"report_token": token, "sync_enabled": True})
            cfg.api_key = APPKEY
            s.add(cfg)
            s.add(Lead(agency_id=agency.id, source_type="manual", status="new",
                       consent_given=True, intent_score=77, topnlab_client_id=555001))
            await s.commit()
            other = Agency(name="Чужое агентство", base_city="Анапа")
            s.add(other)
            await s.flush()
            s.add(Lead(agency_id=other.id, source_type="manual", status="new",
                       consent_given=True, topnlab_client_id=555002))
            await s.commit()
    finally:
        await engine.dispose()

    client = _client()
    body = {"user": {"id": 1, "email": "agent@agency.ru"}, "ids": [555001, 555002, "x"],
            "report_id": 7}
    assert client.post("/api/topnlab/report-webhook?token=wrong-token-0000000", json=body
                       ).status_code == 404
    resp = client.post(f"/api/topnlab/report-webhook?token={token}", json=body)
    assert resp.status_code == 200
    assert resp.headers["access-control-allow-origin"] == "*"
    url = resp.json()["url"]
    assert url.startswith("https://test.local/api/topnlab/report-files/")

    file_resp = client.get(url.replace("https://test.local", ""))
    assert file_resp.status_code == 200
    assert file_resp.headers["access-control-allow-origin"] == "*"
    from io import BytesIO

    from openpyxl import load_workbook

    ws = load_workbook(BytesIO(file_resp.content)).active
    ids = [row[0].value for row in ws.iter_rows(min_row=2)]
    assert ids == [555001]  # the other agency's lead stays out

    events = client.post(f"/api/topnlab/incoming-webhook?token={token}",
                         data={"id": "123", "type": "order"})
    assert events.status_code == 200 and events.json()["received"] is True
    assert client.post("/api/topnlab/incoming-webhook", data={"id": "1"}).status_code == 403


@pytestmark_db
@pytest.mark.asyncio
async def test_owner_settings_switch_the_agency_to_topnlab():
    from tests.helpers import unique_telegram_id

    from app.database import async_session, engine, run_migrations
    from app.dependencies import CurrentManager
    from app.exceptions import AppException
    from app.models.agency import Agency
    from app.models.agency_crm_config import AgencyCRMConfig
    from app.models.manager import Manager
    from app.routers.topnlab import TopnlabSettingsRequest, get_settings, save_settings

    await run_migrations()
    try:
        async with async_session() as s:
            agency = Agency(name="Настройки TopNLab", base_city="Геленджик")
            s.add(agency)
            await s.flush()
            owner = Manager(agency_id=agency.id, name="Владелец", role="owner",
                            telegram_id=unique_telegram_id(), is_active=True)
            staff = Manager(agency_id=agency.id, name="Менеджер", role="manager",
                            telegram_id=unique_telegram_id(), is_active=True)
            old = AgencyCRMConfig(agency_id=agency.id, crm_type="amocrm", is_active=True,
                                  base_url="https://x.amocrm.ru")
            s.add_all([owner, staff, old])
            await s.commit()
            ctx = CurrentManager(manager_id=str(owner.id), agency_id=str(agency.id))
            staff_ctx = CurrentManager(manager_id=str(staff.id), agency_id=str(agency.id))

        async with async_session() as s:
            with pytest.raises(AppException):
                await save_settings(TopnlabSettingsRequest(sync_enabled=True), current=ctx, session=s)

        async with async_session() as s:
            saved = await save_settings(TopnlabSettingsRequest(
                appkey=APPKEY, company_id="207413", virtual_number="+7 861 234-56-78",
                manager_email="agent@agency.ru", sync_enabled=True), current=ctx, session=s)
        assert saved["active"] and saved["has_key"] and saved["sync_enabled"]
        assert saved["virtual_number"] == "8612345678"
        assert saved["incoming_webhook_url"].startswith("https://test.local/api/topnlab/")
        assert APPKEY not in json.dumps(saved)

        async with async_session() as s:
            rows = (await s.execute(__import__("sqlalchemy").select(AgencyCRMConfig).where(
                AgencyCRMConfig.agency_id == agency.id))).scalars().all()
            assert [r.crm_type for r in rows] == ["topnlab"]  # one CRM per agency
            assert rows[0].base_url is None and rows[0].api_key == APPKEY

        async with async_session() as s:
            with pytest.raises(AppException) as err:
                await get_settings(current=staff_ctx, session=s)
            assert err.value.status_code == 403

        async with async_session() as s:
            with pytest.raises(AppException) as err:
                await save_settings(TopnlabSettingsRequest(company_id="abc"), current=ctx, session=s)
            assert err.value.status_code == 400
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_object_behind_an_avito_listing(api):
    """Priority 7: avitoData. Found -> realty and agent; not found -> None."""
    fake = api({"/call/main/avitoData/": lambda r: httpx.Response(200, json=(
        {"status": "ok", "result": {"realty": {"id": 123, "action": "sale"},
                                    "user": {"id": 9, "email": "a@agency.ru"}}}
        if r.url.params["publication_id"] == "112233" else
        {"status": "error", "errors": {"publication_id": "Не нашлось такого размещения"}}))})
    found = await tl.get_object_by_avito_id(_settings(), 112233)
    assert found["realty"]["id"] == 123 and found["user"]["email"] == "a@agency.ru"
    assert fake.calls[0][2] == {"publication_id": "112233", "key": APPKEY}
    assert await tl.get_object_by_avito_id(_settings(), 1) is None
